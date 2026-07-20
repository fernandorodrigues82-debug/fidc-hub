"""Exporta o resultado da simulação para Excel (.xlsx) — várias abas,
formatação profissional, pronto para o analista continuar fora do app.

Roda dentro do próprio Streamlit (openpyxl é dependência pura Python).
Os valores são estáticos (não fórmulas): isto é uma exportação de um
resultado já calculado, não uma planilha-modelo para o usuário editar
premissas — as premissas continuam sendo editadas no app.
"""

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COR_HEADER = "0B5563"
FONTE = "Calibri"

FMT_RS = 'R$ #,##0.00;[RED]-R$ #,##0.00;"-"'
FMT_PCT = "0.00%"
FMT_PCT1 = "0.0%"
FMT_X = '0.00"x"'


def _escrever_df(ws, df: pd.DataFrame, start_row=1, formatos=None,
                 titulo=None):
    formatos = formatos or {}
    row = start_row
    if titulo:
        c = ws.cell(row=row, column=1, value=titulo)
        c.font = Font(name=FONTE, bold=True, size=13, color=COR_HEADER)
        row += 2
    header_row = row
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=header_row, column=j, value=str(col))
        c.font = Font(name=FONTE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COR_HEADER)
        c.alignment = Alignment(horizontal="center")
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(df.columns, start=1):
            val = r[col]
            if pd.isna(val):
                val = None
            elif hasattr(val, "item"):
                val = val.item()
            cell = ws.cell(row=header_row + i, column=j, value=val)
            cell.font = Font(name=FONTE)
            fmt = formatos.get(col)
            if fmt and val is not None:
                cell.number_format = fmt
    for j, col in enumerate(df.columns, start=1):
        maxlen = max([len(str(col))] +
                    [len(str(v)) for v in df[col]]) if len(df) else len(str(col))
        ws.column_dimensions[get_column_letter(j)].width = min(
            45, max(11, maxlen + 2))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    return header_row + len(df) + 2


def gerar_excel(*, nome_fundo: str, originador: dict, estrutura, resumo: dict,
                por_classe: pd.DataFrame, fluxo: pd.DataFrame, suporte: dict,
                ratings=None, mc_stats=None, calibracao: dict | None = None,
                ) -> bytes:
    wb = Workbook()

    # ---------------------------------------------------------- Resumo
    ws0 = wb.active
    ws0.title = "Resumo"
    integra = resumo.get("todas_integras", resumo.get("senior_integra"))
    linhas = [
        ("Fundo", nome_fundo),
        ("Originador", originador.get("razao_social", "—")),
        ("PL total (R$)", estrutura.pl_total),
        ("Taxa de cessão (% a.m.)", estrutura.taxa_cessao_am),
        ("Prazo médio (meses)", estrutura.prazo_medio_meses),
        ("Revolvência (meses)", estrutura.meses_revolvencia),
        ("Inadimplência base (% a.m.)", estrutura.inadimplencia_am),
        ("Stress aplicado (x)", estrutura.stress),
        ("Gatilho de subordinação mínima (%)",
         estrutura.sub_minima if estrutura.sub_minima else None),
        ("Todas as classes íntegras no cenário", "Sim" if integra else "NÃO"),
        ("Perdas totais no cenário (R$)", resumo.get("perdas_totais")),
        ("Break-even da sênior (x inadimplência base)",
         suporte.get("breakeven_mult")),
        ("Perda máxima absorvível (R$)", suporte.get("perda_maxima")),
        ("Perda máxima absorvível (% do PL)",
         suporte.get("perda_maxima_pct_pl")),
    ]
    df_resumo = pd.DataFrame(linhas, columns=["Item", "Valor"])
    _escrever_df(ws0, df_resumo, titulo=f"Memorando — {nome_fundo}",
                formatos={})
    # formatação linha a linha (tipos variados na coluna Valor)
    fmt_por_item = {
        "PL total (R$)": FMT_RS, "Perdas totais no cenário (R$)": FMT_RS,
        "Perda máxima absorvível (R$)": FMT_RS,
        "Taxa de cessão (% a.m.)": FMT_PCT,
        "Inadimplência base (% a.m.)": FMT_PCT,
        "Gatilho de subordinação mínima (%)": FMT_PCT1,
        "Perda máxima absorvível (% do PL)": FMT_PCT1,
        "Stress aplicado (x)": FMT_X,
        "Break-even da sênior (x inadimplência base)": FMT_X,
    }
    for i, (item, _) in enumerate(linhas, start=4):
        if item in fmt_por_item:
            ws0.cell(row=i, column=2).number_format = fmt_por_item[item]

    # ------------------------------------------------------- Por classe
    ws1 = wb.create_sheet("Por classe")
    tabela = por_classe.copy()
    if ratings is not None and len(ratings):
        tabela = tabela.merge(ratings[["classe", "nota"]], on="classe",
                              how="left")
    tabela["integra"] = tabela["integra"].map({True: "Sim", False: "NÃO"})
    _escrever_df(ws1, tabela, formatos={
        "pct": FMT_PCT1, "aporte": FMT_RS, "recebido": FMT_RS,
        "shortfall": FMT_RS, "tir_aa": FMT_PCT})

    # ------------------------------------------------------ Fluxo mensal
    ws2 = wb.create_sheet("Fluxo mensal")
    formatos_fluxo = {c: FMT_RS for c in fluxo.columns
                      if c not in ("mes", "fase", "indice_subordinacao")}
    if "indice_subordinacao" in fluxo.columns:
        formatos_fluxo["indice_subordinacao"] = FMT_PCT1
    _escrever_df(ws2, fluxo, formatos=formatos_fluxo)

    # ------------------------------------------------------- Monte Carlo
    if mc_stats is not None and len(mc_stats):
        ws3 = wb.create_sheet("Monte Carlo")
        mc = mc_stats.copy()
        if ratings is not None and len(ratings):
            mc = mc.merge(ratings[["classe", "nota"]], on="classe",
                          how="left")
        _escrever_df(ws3, mc, formatos={
            "tir_media": FMT_PCT, "tir_p5": FMT_PCT, "tir_p50": FMT_PCT,
            "tir_p95": FMT_PCT, "prob_nao_integral": FMT_PCT,
            "prob_perda_principal": FMT_PCT})

    # ------------------------------------------------------- Calibração
    if calibracao:
        ws4 = wb.create_sheet("Calibração carteira")
        linhas_c = [
            ("Meses de histórico", calibracao["n_meses"]),
            ("Período coberto", calibracao.get("periodo", "—")),
            ("Taxa média observada (% a.m.)", calibracao["taxa_media_am"]),
            ("Volatilidade calibrada", calibracao["vol"]),
            ("Persistência (ρ) calibrada", calibracao["rho"]),
        ]
        df_c = pd.DataFrame(linhas_c, columns=["Item", "Valor"])
        _escrever_df(ws4, df_c, titulo="Calibração da carteira real")
        ws4.cell(row=6, column=2).number_format = FMT_PCT

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
