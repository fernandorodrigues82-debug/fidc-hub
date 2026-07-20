"""Gera o memorando de comitê (.docx) a partir de uma simulação já rodada.

Roda dentro do próprio app Streamlit (python-docx é dependência pura Python,
funciona igual em produção). Recebe apenas dados já calculados — não
reexecuta motor nenhum — para o documento refletir exatamente o que o
usuário viu na tela.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

COR_TITULO = RGBColor(0x0B, 0x55, 0x63)


class _Numerador:
    """Contador de seções — evita numeração manual frágil e duplicada."""
    def __init__(self):
        self.n = 0

    def titulo(self, doc, texto, nivel=1):
        self.n += 1
        h = doc.add_heading(level=nivel)
        run = h.add_run(f"{self.n}. {texto}")
        run.font.color.rgb = COR_TITULO
        return h


def _tabela(doc, cabecalho, linhas, estilo="Light Grid Accent 1"):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = estilo
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, texto in enumerate(cabecalho):
        cell = t.rows[0].cells[i]
        cell.text = str(texto)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for linha in linhas:
        row = t.add_row().cells
        for i, valor in enumerate(linha):
            row[i].text = str(valor)
    return t


def _p(doc, texto, tamanho=10, italic=False, bold=False, cor=None):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.font.size = Pt(tamanho)
    run.italic = italic
    run.bold = bold
    if cor:
        run.font.color.rgb = cor
    return par


def gerar_memorando(*, nome_fundo: str, originador: dict, estrutura,
                    resumo: dict, por_classe, suporte: dict,
                    calibracao: dict | None = None,
                    mc_stats=None, responsavel: str = "") -> bytes:
    """Monta o memorando e devolve os bytes do .docx (pronto para
    st.download_button)."""
    doc = Document()
    num = _Numerador()
    for style_name in ("Normal",):
        doc.styles[style_name].font.name = "Calibri"
        doc.styles[style_name].font.size = Pt(10.5)

    # --------------------------------------------------------- capa/cabeçalho
    h = doc.add_heading(level=0)
    r = h.add_run(f"Memorando de Comitê — {nome_fundo}")
    r.font.color.rgb = COR_TITULO
    _p(doc, f"Originador: {originador.get('razao_social', '—')}"
           f"  ·  CNPJ: {originador.get('cnpj') or '—'}", tamanho=10.5)
    _p(doc, f"Setor: {originador.get('setor') or '—'}  ·  "
           f"Recebível: {originador.get('tipo_recebivel') or '—'}",
       tamanho=10.5)
    _p(doc, f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
           + (f"  ·  Responsável: {responsavel}" if responsavel else ""),
       tamanho=9, italic=True)
    doc.add_paragraph()

    # --------------------------------------------------------- 1. sumário
    num.titulo(doc, "Sumário executivo")
    integra = resumo.get("todas_integras", resumo.get("senior_integra"))
    veredicto = ("Estrutura aprovável — todas as classes íntegras no "
                "cenário simulado." if integra else
                "ATENÇÃO — há classe(s) com perda no cenário simulado; "
                "revisar antes de submeter ao comitê.")
    _p(doc, veredicto, bold=True,
       cor=RGBColor(0x1C, 0x6B, 0x2E) if integra else RGBColor(0xB3, 0x3A, 0x3A))
    if originador.get("score") is not None:
        _p(doc, f"Score de estruturabilidade do originador: "
               f"{originador['score']}/100. Inadimplência histórica "
               f"declarada: {originador.get('inadimplencia_hist', '—')}%. "
               f"Concentração top-10 sacados: "
               f"{originador.get('concentracao_top10', '—')}%.")
    if suporte.get("breakeven_mult"):
        be = suporte["breakeven_mult"]
        be_txt = f"≥ {be:.0f}x" if be >= 29.9 else f"{be:.1f}x"
        _p(doc, f"A estrutura suporta {be_txt} a inadimplência base "
               f"antes de a classe mais sênior sofrer perda, absorvendo "
               f"até R$ {suporte['perda_maxima']/1e6:,.1f} milhões em "
               f"créditos "
               f"({suporte['perda_maxima_pct_pl']*100:.1f}% do PL).")

    # --------------------------------------------------------- 2. estrutura
    num.titulo(doc, "Estrutura de capital proposta")
    _p(doc, f"PL total: R$ {estrutura.pl_total/1e6:,.1f} milhões  ·  "
           f"Taxa de cessão: {estrutura.taxa_cessao_am*100:.2f}% a.m.  ·  "
           f"Prazo médio dos recebíveis: {estrutura.prazo_medio_meses} "
           f"mês(es)  ·  Revolvência: {estrutura.meses_revolvencia} meses")
    if estrutura.sub_minima:
        _p(doc, f"Gatilho de subordinação mínima (evento de avaliação): "
               f"{estrutura.sub_minima*100:.0f}%.")
    cab = ["Classe", "% do PL", "Taxa-alvo (a.m.)", "Ordem de senioridade"]
    linhas = [[c.nome, f"{c.pct*100:.1f}%",
              "residual" if c.residual else f"{c.taxa_am*100:.2f}%",
              i + 1] for i, c in enumerate(estrutura.classes)]
    _tabela(doc, cab, linhas)
    doc.add_paragraph()

    # ------------------------------------------------- 3. resultado simulado
    num.titulo(doc, "Resultado da simulação (cenário determinístico)")
    cab = ["Classe", "Aporte (R$ mi)", "Recebido (R$ mi)", "TIR (a.a.)",
          "Íntegra"]
    linhas = []
    for _, row in por_classe.iterrows():
        tir = row.get("tir_aa")
        tir_txt = f"{tir*100:.2f}%" if tir is not None else "—"
        linhas.append([row["classe"], f"{row['aporte']/1e6:,.1f}",
                       f"{row['recebido']/1e6:,.1f}", tir_txt,
                       "Sim" if row["integra"] else "NÃO"])
    _tabela(doc, cab, linhas)
    doc.add_paragraph()
    _p(doc, f"Perdas totais no cenário: R$ "
           f"{resumo.get('perdas_totais', 0)/1e6:,.1f} milhões "
           f"({resumo.get('perdas_vs_sub', 0)*100:.0f}% do colchão de "
           f"subordinação inicial).")
    if resumo.get("gatilho_mes"):
        _p(doc, f"⚡ O gatilho de subordinação mínima foi acionado no mês "
               f"{resumo['gatilho_mes']} deste cenário: a revolvência foi "
               f"interrompida e a estrutura passou a amortizar "
               f"antecipadamente por senioridade.")

    # --------------------------------------------------- 4. suporte/stress
    num.titulo(doc, "Suporte da estrutura — teste de stress")
    if suporte.get("breakeven_mult") is None:
        _p(doc, "A classe mais sênior já sofre perda no cenário base — "
               "não há colchão de stress a reportar.",
           cor=RGBColor(0xB3, 0x3A, 0x3A), bold=True)
    else:
        be = suporte["breakeven_mult"]
        be_txt = f"≥ {be:.0f}x" if be >= 29.9 else f"{be:.1f}x"
        cab = ["Métrica", "Valor"]
        linhas = [
            ["Suporta até (múltiplo da inadimplência base)", be_txt],
            ["Inadimplência mensal equivalente no limite",
             f"{suporte['inad_am_equivalente']*100:.2f}% a.m."],
            ["Perda absorvida no limite", f"R$ "
             f"{suporte['perda_maxima']/1e6:,.1f} mi"],
            ["Perda no limite / PL",
             f"{suporte['perda_maxima_pct_pl']*100:.1f}%"],
        ]
        _tabela(doc, cab, linhas, estilo="Light List Accent 1")

    # ------------------------------------------- 5. distribuição de retornos
    if mc_stats is not None and len(mc_stats):
        num.titulo(doc, "Distribuição de retornos (Monte Carlo)")
        _p(doc, "Simulação estocástica da inadimplência mensal com "
               "persistência de ciclo de crédito, em torno da "
               "inadimplência base e do stress definidos na estrutura "
               "acima.")
        cab = ["Classe", "TIR média", "TIR p5 (pior 5%)", "TIR mediana",
              "TIR p95", "Prob. não íntegra", "Prob. perder principal"]
        linhas = []
        for _, row in mc_stats.iterrows():
            def _pct(v):
                return f"{v*100:.2f}%" if v == v else "—"  # NaN-safe
            linhas.append([
                row["classe"], _pct(row["tir_media"]), _pct(row["tir_p5"]),
                _pct(row["tir_p50"]), _pct(row["tir_p95"]),
                _pct(row["prob_nao_integral"]),
                _pct(row["prob_perda_principal"])])
        _tabela(doc, cab, linhas)
    else:
        num.titulo(doc, "Distribuição de retornos (Monte Carlo)")
        _p(doc, "Não executada para esta estrutura neste memorando. "
               "Disponível na página 'Simulação de Retornos' do FIDC Hub.",
           italic=True)

    # --------------------------------------------------- calibração (se houver)
    if calibracao:
        num.titulo(doc, "Calibração da carteira real")
        _p(doc, f"Volatilidade e persistência do ciclo de crédito "
               f"calibradas a partir de {calibracao['n_meses']} meses de "
               f"histórico de vencimentos maturados "
               f"({calibracao.get('periodo', '—')}): volatilidade "
               f"{calibracao['vol']}, persistência (ρ) {calibracao['rho']}, "
               f"taxa média observada "
               f"{calibracao['taxa_media_am']*100:.2f}% a.m.")
        _p(doc, "Metodologia: proxy de aging por safra de vencimento "
               "(% do valor 90+ dias em atraso na data do relatório) — "
               "calibra a forma da distribuição (oscilação e persistência), "
               "não substitui a due diligence para o nível da "
               "inadimplência base.", italic=True, tamanho=9)

    # --------------------------------------------------------- observações
    if originador.get("notas"):
        num.titulo(doc, "Notas de due diligence")
        _p(doc, originador["notas"])

    doc.add_paragraph()
    _p(doc, "Documento gerado automaticamente pelo FIDC Hub a partir dos "
           "parâmetros simulados. Ferramenta de apoio à decisão — não "
           "substitui assessoria jurídica nem a validação regulatória "
           "(CVM/ANBIMA) da estrutura.", tamanho=8.5, italic=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
